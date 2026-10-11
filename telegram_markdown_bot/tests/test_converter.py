from pathlib import Path

import pymupdf
from docx import Document
from docx.shared import Pt
from openpyxl import Workbook
from pptx import Presentation
from PIL import Image, ImageDraw

from converter import clean_compact, convert_file, safe_table, estimate_tokens


def test_clean_keeps_arabic_math():
    src = 'القوى والأسس\n\n\n2^3 = 2 × 2 × 2\nPage 3 of 7\n\nتمارين 1|2\n'
    out = clean_compact(src)
    assert '2^3 = 2 × 2 × 2' in out
    assert 'Page 3 of 7' not in out
    assert 'تمارين 1|2' in out


def test_safe_table_pipe():
    table = safe_table([['سؤال', '1|2'], ['قيمة', '7']])
    assert r'1\|2' in table
    assert table.count('\n') == 3


def test_txt(tmp_path):
    p = tmp_path / 'arabic.txt'
    p.write_text('هذا نص عربي\n\n\nالقوى ٢^٣\n', encoding='utf8')
    raw = convert_file(p)
    compact = convert_file(p, mode='compact')
    assert 'هذا نص عربي' in raw.markdown
    assert 'القوى ٢^٣' in compact.markdown
    assert len(compact.markdown) < len(raw.markdown)


def test_docx_and_table(tmp_path):
    p = tmp_path / 'lesson.docx'
    doc = Document()
    doc.add_heading('القوى والأسس', 1)
    doc.add_paragraph('اكتب 2^4 على صورة ضرب العامل في نفسه.')
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = 'أ'
    table.cell(0, 1).text = 'ب'
    table.cell(1, 0).text = '2'
    table.cell(1, 1).text = '4'
    doc.save(p)
    md = convert_file(p).markdown
    assert '# القوى والأسس' in md
    assert '2^4' in md
    assert '| 2 | 4 |' in md


def test_xlsx(tmp_path):
    p = tmp_path / 'data.xlsx'
    wb = Workbook()
    ws = wb.active
    ws.title = 'درجات'
    ws.append(['طالب', 'درجة'])
    ws.append(['أحمد', 94])
    wb.save(p)
    result = convert_file(p).markdown
    assert 'درجات' in result
    assert 'أحمد' in result and '94' in result


def test_pptx(tmp_path):
    p = tmp_path / 'deck.pptx'
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = 'Lesson One'
    slide.placeholders[1].text = 'Formula: 2^3'
    prs.save(p)
    result = convert_file(p).markdown
    assert 'Lesson One' in result and 'Formula: 2^3' in result


def test_pdf(tmp_path):
    p = tmp_path / 'text.pdf'
    doc = pymupdf.open()
    pg = doc.new_page()
    pg.insert_text((72, 80), 'Math exercises - powers 2^3 = 8')
    doc.save(p)
    doc.close()
    result = convert_file(p).markdown
    assert 'Math exercises' in result


def test_html(tmp_path):
    p = tmp_path / 'page.html'
    p.write_text('<html><body><h1>عنوان</h1><script>bad()</script><p>نص تجريبي</p></body></html>', encoding='utf8')
    md = convert_file(p).markdown
    assert '# عنوان' in md
    assert 'نص تجريبي' in md
    assert 'bad()' not in md


def test_tokens():
    n, method = estimate_tokens('أهلًا world')
    assert n > 0
