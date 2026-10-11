"""Local document-to-Markdown conversion. No LLM or paid APIs are used.

Improved PDF extraction for 6-panel Arabic handouts, preserving superscripts.
"""
from __future__ import annotations

import csv
import io
import re
import unicodedata
import zipfile
from dataclasses import dataclass, field
from pathlib import Path

import pymupdf
from bs4 import BeautifulSoup
from docx import Document
from docx.table import Table as DocxTable
from docx.text.paragraph import Paragraph as DocxParagraph
from markdownify import markdownify
from openpyxl import load_workbook
from PIL import Image, ImageOps, UnidentifiedImageError
from pptx import Presentation

SUPPORTED = {'.pdf', '.docx', '.pptx', '.xlsx', '.txt', '.md', '.csv', '.html', '.htm', '.png', '.jpg', '.jpeg', '.webp'}
MAX_PDF_PAGES = 40
MAX_OFFICE_UNCOMPRESSED = 100 * 1024 * 1024
MAX_OUTPUT_CHARS = 8_000_000
Image.MAX_IMAGE_PIXELS = 28_000_000

class ConversionError(Exception):
    """Conversion issue safe to show to the Telegram user."""

@dataclass
class ConversionResult:
    markdown: str
    raw_markdown: str
    warnings: list[str] = field(default_factory=list)


def safe_table(rows: list[list[str]]) -> str:
    if not rows:
        return ''
    width = max(len(row) for row in rows)
    if not width:
        return ''
    def cell(v: str) -> str:
        return str(v).replace('|', r'\|').replace('\r', '').replace('\n', ' / ').strip()
    padded = [[cell(row[i]) if i < len(row) else '' for i in range(width)] for row in rows]
    header = [f'عمود {i+1}' for i in range(width)]
    return '\n'.join(['| ' + ' | '.join(header) + ' |', '| ' + ' | '.join(['---']*width) + ' |'] + [
        '| ' + ' | '.join(row) + ' |' for row in padded])


def assert_safe_office_zip(path: Path) -> None:
    try:
        with zipfile.ZipFile(path) as archive:
            infos = archive.infolist()
            if len(infos) > 2500 or sum(i.file_size for i in infos) > MAX_OFFICE_UNCOMPRESSED:
                raise ConversionError('ملف Office كبير جدًا بعد فك الضغط؛ لا يمكن معالجته بأمان.')
            if any(i.file_size > 30*1024*1024 for i in infos):
                raise ConversionError('يحتوي الملف على جزء داخلي كبير جدًا.')
    except zipfile.BadZipFile as ex:
        raise ConversionError('ملف Office تالف أو امتداده غير صحيح.') from ex


def clean_compact(markdown: str) -> str:
    out = []
    for line in markdown.replace('\r\n','\n').replace('\r','\n').split('\n'):
        line=line.rstrip()
        if re.fullmatch(r'\s*[-_=*•·.]{5,}\s*',line):
            continue
        if re.fullmatch(r'\s*(?:Page\s+\d+(?:\s+of\s+\d+)?|صفحة\s+\d+(?:\s+من\s+\d+)?)\s*',line,re.I):
            continue
        if not line and out and not out[-1]:
            continue
        out.append(line)
    return '\n'.join(out).strip()+'\n'


def _math_line(line: dict) -> str | None:
    """Rebuild positioned math, distinguishing smaller adjacent exponent glyphs.

    Returns None for ordinary text, where the PDF's native RTL order is better.
    """
    spans = line.get('spans', [])
    if not spans:
        return None
    all_text = ''.join(s['text'] for s in spans)
    has_ops = any(ch in all_text for ch in '=×÷')
    numeric_only = len(spans) >= 2 and all(
        ch.isdigit() or ch.isspace() for s in spans for ch in s['text'])
    if not (has_ops or numeric_only):
        return None
    parts = []
    prior = None
    for span in sorted(spans, key=lambda s:s['bbox'][0]):
        text=span['text']
        if re.fullmatch(r'[\u0660-\u0669\u06f0-\u06f90-9]+',text):
            # Arabic numeral sequences in this font are stored opposite to display.
            if len(text)>1:
                text=text[::-1]
            text=''.join(str(unicodedata.digit(ch)) if ch.isdigit() else ch for ch in text)
            if (prior and prior['text'].strip().isdigit()
                    and span['size'] < prior['size'] * .82
                    and abs(span['bbox'][0]-prior['bbox'][2]) < 4):
                parts.append('^'+text)
            else:
                parts.append(text)
        else:
            parts.append(text)
        prior=span
    return ''.join(parts).strip()


def _looks_like_six_panel_handout(page: pymupdf.Page) -> bool:
    """Conservative detection: distinct three columns AND a full horizontal fold gap."""
    w,h=page.rect.width,page.rect.height
    if not (1.32 < h/w < 1.52):
        return False
    blocks=[b for b in page.get_text('dict')['blocks'] if b.get('type')==0]
    if len(blocks)<35:
        return False
    occupied=[0]*6
    top_bottom=[]
    bottom_top=[]
    for block in blocks:
        x0,y0,x1,y1=block['bbox']
        xc=(x0+x1)/2; yc=(y0+y1)/2
        c=min(2,max(0,int(xc/(w/3))))
        r=int(yc>=h/2)
        occupied[r*3+c]+=1
        if r==0: top_bottom.append(y1)
        else: bottom_top.append(y0)
    if not all(count>=4 for count in occupied):
        return False
    if not top_bottom or not bottom_top:
        return False
    # A three-column newspaper does not have a gap across the whole page.
    return (min(bottom_top)-max(top_bottom)) >= .025*h


def _six_panel_markdown(page: pymupdf.Page) -> str:
    """Read separate brochure panels rather than interleave words across columns."""
    cells={i:[] for i in range(6)}
    w,h=page.rect.width,page.rect.height
    for block in page.get_text('dict')['blocks']:
        if block.get('type') != 0:
            continue
        x0,y0,x1,y1=block['bbox']
        c=min(2,max(0,int(((x0+x1)/2)/(w/3))))
        r=int(((y0+y1)/2)>=h/2)
        cell=r*3+c
        for line in block['lines']:
            txt=_math_line(line)
            if txt is None:
                txt=''.join(s['text'] for s in line.get('spans', [])).strip()
            if txt:
                cells[cell].append((line['bbox'][1],line['bbox'][0],txt))

    def normalize(txt: str) -> str:
        # Tolerate faulty Arabic font-to-Unicode mappings in this document type.
        for a,b in [('األ','الأ'),('إىل','إلى'),('عىل','على'),
                    ('أوˀ','أولًا'),('ثانياًلا','ثانيًا'),('ثالثاًلا','ثالثًا'),
                    ('عامˁً','عامًا')]:
            txt=txt.replace(a,b)
        txt=txt.replace('.','.')
        txt=re.sub(r'^(العوامل|القوة|الأساس|الأُس|الأسّ|الصيغة الأسية|الصيغة القياسية)\.',r'\1: ',txt)
        txt=re.sub(r'^=([0-9]+\^[0-9]+)$',r'\1 =',txt)
        txt=re.sub(r'^=([0-9]+)$',r'= \1',txt)
        return txt
    pieces=[]
    for seq,idx in enumerate([2,1,0,5,4,3],start=1):
        pieces.append(f'## القسم {seq}')
        ordered=sorted(cells[idx],key=lambda q:(q[0],-q[1]))
        groups=[]
        for y,x,txt in ordered:
            if groups and abs(y-groups[-1][0])<=6.5:
                groups[-1][1].append((x,txt))
            else:
                groups.append([y,[(x,txt)]])
        last_y=None
        for y,frags in groups:
            if last_y is not None and y-last_y>28:
                pieces.append('')
            frags.sort(key=lambda it:-it[0])
            txt=' '.join(part for _,part in frags)
            txt=normalize(txt)
            pieces.append(txt)
            last_y=y
        pieces.append('')
    return '\n'.join(pieces)


def _pixmap_image(page) -> Image.Image:
    pix=page.get_pixmap(matrix=pymupdf.Matrix(1.9,1.9),alpha=False)
    return Image.open(io.BytesIO(pix.tobytes('png')))


def _image_ocr(img: Image.Image) -> str:
    try:
        import pytesseract
        img=ImageOps.exif_transpose(img).convert('RGB')
        return pytesseract.image_to_string(img,lang='ara+eng',config='--psm 3').strip()
    except (ImportError, OSError) as ex:
        raise ConversionError('التعرف على النص من الصور غير متاح. ثبّت Tesseract مع اللغتين ara وeng.') from ex
    except Exception as ex:
        raise ConversionError('فشلت قراءة النص من الصورة؛ تحقق من وضوحها.') from ex


def _pdf(path: Path, *,force_ocr: bool) -> tuple[str,list[str]]:
    warnings=[]
    try:
        with pymupdf.open(str(path)) as doc:
            if doc.needs_pass:
                raise ConversionError('الملف محمي بكلمة مرور. أرسل نسخة غير محمية.')
            count=doc.page_count
            if not count:
                raise ConversionError('ملف PDF لا يحتوي على صفحات.')
            if count>MAX_PDF_PAGES:
                raise ConversionError(f'هذا الملف فيه {count} صفحة. الحد الحالي {MAX_PDF_PAGES} صفحة.')
            lengths=[len(page.get_text().strip()) for page in doc]
            scanned_pages=sum(length<10 for length in lengths)
            if scanned_pages:
                warnings.append(f'يوجد {scanned_pages} صفحة دون نص رقمي واضح؛ تعتمد دقة ناتجها على OCR.')
            # Preserve semantics and equations in a 3-column x 2-row printout.
            # Force OCR overrides this because the text layer may be corrupted.
            if count==1 and not force_ocr and _looks_like_six_panel_handout(doc[0]):
                text=_six_panel_markdown(doc[0])
                warnings.append('تم تقسيم المطوية إلى ستة أقسام بحسب مواضعها، وتحليل الأسس من حجم الحروف. راجع الصيغ المعقدة.')
                return text, warnings
        try:
            import pymupdf4llm
        except ImportError:
            pymupdf4llm=None
        if pymupdf4llm is not None:
            text=pymupdf4llm.to_markdown(
                str(path),ocr_language='ara+eng',force_ocr=force_ocr,
                use_ocr=True,write_images=False,embed_images=False,
                page_separators=False,show_progress=False)
        else:
            with pymupdf.open(str(path)) as doc:
                parts=[]
                for page in doc:
                    s=page.get_text('text',sort=True).strip()
                    if force_ocr or not s:
                        s=_image_ocr(_pixmap_image(page))
                    parts.append(s)
                text='\n\n'.join(parts)
            warnings.append('لم يتوفر محرك الجداول المتقدم؛ استخدم المتطلبات كاملة للحصول على أفضل نتيجة.')
        if not text or not text.strip():
            raise ConversionError('تعذّر استخراج النص. جرّب /ocr on ثم أرسل الملف مرة ثانية.')
        if force_ocr:
            warnings.append('استُخدم وضع OCR الإجباري؛ راجع الرموز الرياضية والأرقام بعناية.')
        warnings.append('الرسومات والصور غير النصية لا تتحول تلقائيًا إلى وصف دقيق.')
        return str(text),warnings
    except ConversionError:
        raise
    except Exception as ex:
        raise ConversionError('تعذرت قراءة PDF. تحقق أن الملف صالح، أو جرّب وضع OCR.') from ex


def _image(path: Path) -> tuple[str,list[str]]:
    try:
        with Image.open(path) as img:
            if img.width*img.height > Image.MAX_IMAGE_PIXELS:
                raise ConversionError('أبعاد الصورة كبيرة جدًا؛ أرسل صورة أصغر.')
            value=_image_ocr(img)
    except (Image.DecompressionBombError,UnidentifiedImageError) as ex:
        raise ConversionError('الصورة تالفة أو عالية الدقة بصورة غير آمنة.') from ex
    if not value:
        raise ConversionError('لم يظهر نص يمكن استخراجه من الصورة.')
    return value,['الصور المعالجة بـOCR قد تحتوي على أخطاء في الأرقام والرموز.']


def _docx(path: Path) -> tuple[str,list[str]]:
    assert_safe_office_zip(path)
    doc=Document(str(path))
    parts=[]
    for block in doc.iter_inner_content():
        if isinstance(block,DocxParagraph):
            value=block.text.strip()
            if not value:continue
            style=block.style.name.lower() if block.style else ''
            if style.startswith('heading'):
                found=re.search(r'\d+',style)
                level=max(1,min(6,int(found.group()) if found else 1))
                parts.append('#'*level+' '+value)
            elif any(s in style for s in ('list bullet','list paragraph','قائمة نقطية')):
                parts.append('- '+value)
            elif 'list number' in style or 'قائمة رقمية' in style:
                parts.append('1. '+value)
            else:parts.append(value)
        elif isinstance(block,DocxTable):
            parts.append(safe_table([[c.text for c in row.cells] for row in block.rows]))
    return '\n\n'.join(parts),['الصور والأشكال المضمّنة في Word لا يتحول معناها تلقائيًا إلى نص.']


def _pptx(path: Path) -> tuple[str,list[str]]:
    assert_safe_office_zip(path)
    presentation=Presentation(str(path))
    parts=[]
    for idx,slide in enumerate(presentation.slides,start=1):
        parts.append(f'## شريحة {idx}')
        for shape in sorted(slide.shapes,key=lambda s:(s.top,s.left)):
            if shape.has_text_frame:
                text='\n'.join(p.text for p in shape.text_frame.paragraphs if p.text.strip()).strip()
                if text:parts.append(text)
            if shape.has_table:
                parts.append(safe_table([[c.text for c in r.cells] for r in shape.table.rows]))
    return '\n\n'.join(parts),['الرسوم البيانية والصور في الشرائح لا يمكن فهمها بالكامل دون نموذج رؤية.']


def _xlsx(path: Path) -> tuple[str,list[str]]:
    assert_safe_office_zip(path)
    wb=load_workbook(str(path),read_only=True,data_only=True)
    try:
        contents=[]
        for ws in wb.worksheets:
            contents.append('# '+ws.title)
            if ws.max_column>75 or (ws.max_row or 0)*(ws.max_column or 0)>120_000:
                raise ConversionError('ورقة Excel كبيرة جدًا للتحويل في النسخة الحالية.')
            rows=[]
            for row in ws.iter_rows(values_only=True):
                values=['' if it is None else str(it) for it in row]
                if any(v.strip() for v in values):rows.append(values)
            if rows:
                width=max(i+1 for row in rows for i,v in enumerate(row) if v.strip())
                contents.append(safe_table([r[:width] for r in rows]))
        return '\n\n'.join(contents),['قد لا تتوفر القيم المحسوبة للصيغ إذا لم يحفظ Excel ناتجها داخل الملف.']
    finally:wb.close()


def _text(path: Path) -> tuple[str,list[str]]:
    data=path.read_bytes()
    codecs=('utf-16',) if data.startswith((b'\xff\xfe',b'\xfe\xff')) else ('utf-8-sig','cp1256','cp1252')
    for enc in codecs:
        try:
            decoded=data.decode(enc)
            if '\x00' in decoded and not data.startswith((b'\xff\xfe',b'\xfe\xff')):continue
            return decoded,[]
        except UnicodeError:pass
    raise ConversionError('تعذر قراءة ترميز الملف النصي.')


def _csv(path: Path) -> tuple[str,list[str]]:
    text,_=_text(path)
    try:
        dialect=csv.Sniffer().sniff(text[:4096],delimiters=',;\t')
        rows=list(csv.reader(io.StringIO(text),dialect))
    except csv.Error:
        rows=list(csv.reader(io.StringIO(text)))
    if len(rows)>15000:raise ConversionError('ملف CSV يحتوي على صفوف أكثر من الحد المسموح.')
    return safe_table(rows),[]


def _html(path: Path) -> tuple[str,list[str]]:
    text,_=_text(path)
    soup=BeautifulSoup(text,'html.parser')
    for element in soup(['script','style','noscript']):element.decompose()
    return markdownify(str(soup.body or soup),heading_style='ATX'),[]


def convert_file(path: str|Path,mode: str='faithful',force_ocr: bool=False) -> ConversionResult:
    path=Path(path)
    suffix=path.suffix.lower()
    if suffix not in SUPPORTED:
        raise ConversionError('نوع الملف غير مدعوم. استخدم PDF أو DOCX أو PPTX أو XLSX أو TXT أو CSV أو HTML أو صورة.')
    if mode not in ('faithful','compact'):
        raise ValueError('Unknown conversion mode')
    handlers={
        '.pdf':lambda:_pdf(path,force_ocr=force_ocr),
        '.docx':lambda:_docx(path),'.pptx':lambda:_pptx(path),
        '.xlsx':lambda:_xlsx(path),'.txt':lambda:_text(path),
        '.md':lambda:_text(path),'.csv':lambda:_csv(path),
        '.html':lambda:_html(path),'.htm':lambda:_html(path),
        '.png':lambda:_image(path),'.jpg':lambda:_image(path),
        '.jpeg':lambda:_image(path),'.webp':lambda:_image(path),
    }
    try:
        raw,warnings=handlers[suffix]()
    except ConversionError:raise
    except Exception as ex:
        raise ConversionError('تعذر تحويل هذا الملف. ربما هو تالف أو لا يحتوي على نص قابل للاستخراج.') from ex
    raw=raw.replace('\r\n','\n').replace('\r','\n').strip()+'\n'
    if not raw.strip():raise ConversionError('الملف لا يحتوي على نص قابل للاستخراج.')
    out=clean_compact(raw) if mode=='compact' else raw
    if len(out)>MAX_OUTPUT_CHARS:
        raise ConversionError('الناتج كبير جدًا؛ قسّم الملف إلى أجزاء أصغر.')
    return ConversionResult(markdown=out,raw_markdown=raw,warnings=warnings)


def estimate_tokens(text: str) -> tuple[int,str]:
    try:
        import tiktoken
        enc=tiktoken.get_encoding('cl100k_base')
        return len(enc.encode(text)),'cl100k_base'
    except Exception:
        return max(1,round(len(text)/3)),'تقريب بالحروف'
